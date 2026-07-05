"""
Tạo bộ file Excel MẪU để thử nghiệm công cụ (KHÔNG phải dữ liệu thật).

Bộ mẫu mô phỏng đúng bố cục các file thực tế:
  - Bảng kê mua vào (01-1/HT): 2 sheet, dữ liệu ở sheet "BKMV", tiêu đề dòng 10-11,
    dữ liệu từ dòng 12.
  - Danh sách DN rủi ro: 3 sheet khác cấu trúc.
  - File HĐĐT: sheet "DanhSach", tiêu đề dòng 4.
  - File TMS (Danh bạ NNT): tiêu đề dòng 2, MST cột 4, trạng thái cột 45, ngày đóng cột 54.

Chạy:  python scripts/make_samples.py   -> kết quả trong thư mục samples/
"""

import os

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(HERE, "samples")
os.makedirs(OUT, exist_ok=True)


def _set(ws, row, col, value):
    ws.cell(row=row, column=col, value=value)


def make_main():
    """Bảng kê mua vào: sheet phụ 'KQ rà MV' + sheet dữ liệu 'BKMV'."""
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "KQ rà MV"
    _set(ws0, 1, 1, "Kết quả rà soát (sheet tóm tắt - không phải dữ liệu)")

    ws = wb.create_sheet("BKMV")
    # Dòng 1-9: khối tiêu đề biểu mẫu
    _set(ws, 2, 1, "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO")
    _set(ws, 5, 1, "[02] Tên người nộp thuế: CÔNG TY MẪU DEMO")
    _set(ws, 6, 1, "[03] Mã số thuế: 0316139015")

    # Dòng 10: nhãn tiêu đề hàng trên
    _set(ws, 10, 1, "STT")
    _set(ws, 10, 9, "Tên người bán")
    _set(ws, 10, 12, "Mã số thuế người bán")
    _set(ws, 10, 13, "Tên hàng hóa, dịch vụ")
    _set(ws, 10, 17, "Giá trị HHDV mua vào chưa có thuế GTGT")
    _set(ws, 10, 19, "Tiền thuế GTGT")
    # Dòng 11: nhãn hàng dưới
    _set(ws, 11, 5, "Ký hiệu")
    _set(ws, 11, 6, "Số")
    _set(ws, 11, 7, "Ngày, tháng, năm")

    # Dòng 12+: dữ liệu (cột: 6=Số, 7=Ngày, 9=Tên, 12=MST, 13=Hàng, 17=chưa thuế, 19=thuế)
    rows = [
        ("0000123", "2022-01-15", "Ngân hàng ABC", "0301224067", "Phí ngân hàng", 5250000, 525000),
        ("0000123", "2022-01-15", "Ngân hàng ABC", "0301224067", "Phí ngân hàng", 5250000, 525000),  # trùng
        ("0000045", "2022-02-20", "Công ty Rượu XYZ", "0312816241", "Chai rượu vang biếu khách", 5000000, 500000),
        ("0000046", "2022-03-01", "Công ty Golf", "0313025323", "Bộ gậy chơi golf", 2000000, 200000),
        # hóa đơn 2 dòng (cùng số/MST/ngày) -> tổng thuế = 300000 khớp HĐĐT
        ("0000789", "2022-04-10", "Công ty Vật tư", "0400111222", "Thép hộp", 2000000, 200000),
        ("0000789", "2022-04-10", "Công ty Vật tư", "0400111222", "Ốc vít", 1000000, 100000),
        ("0000790", "2022-06-15", "Công ty Quà", "0400111222", "Quà tặng khách hàng", 3000000, 300000),
    ]
    r = 12
    for sohd, ngay, ten, mst, hang, pretax, vat in rows:
        _set(ws, r, 1, r - 11)
        _set(ws, r, 6, sohd)
        _set(ws, r, 7, ngay)
        _set(ws, r, 9, ten)
        _set(ws, r, 12, mst)
        _set(ws, r, 13, hang)
        _set(ws, r, 17, pretax)
        _set(ws, r, 19, vat)
        r += 1

    wb.save(os.path.join(OUT, "01_du_lieu_chinh.xlsx"))


def make_tms():
    """Danh bạ NNT (mẫu): tiêu đề dòng 2, MST cột 4, trạng thái cột 45, ngày đóng cột 54."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NNT"
    ncols = 76
    _set(ws, 1, 1, "DANH BẠ NGƯỜI NỘP THUẾ")
    # dòng 2: tiêu đề
    _set(ws, 2, 4, "Mã số thuế")
    _set(ws, 2, 45, "Trạng thái ĐKT tổ chức")
    _set(ws, 2, 54, "Ngày đóng trạng thái tổ chức")
    for j in range(1, ncols + 1):
        if ws.cell(row=2, column=j).value is None:
            _set(ws, 2, j, f"Cột {j}")
    # dữ liệu từ dòng 3
    data = [
        ("0301224067", "NNT đang hoạt động", ""),
        ("0312816241", "NNT ngừng hoạt động", "2022-02-10"),
        ("0400111222", "NNT đang hoạt động", ""),
    ]
    for i, (mst, tt, ngay) in enumerate(data, start=3):
        _set(ws, i, 4, mst)
        _set(ws, i, 45, tt)
        _set(ws, i, 54, ngay)
    wb.save(os.path.join(OUT, "02_tms.xlsx"))


def make_risk():
    """Danh sách DN rủi ro — 3 sheet khác cấu trúc."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Cơ quan điều tra - UPDATE"
    hdr = ["STT", "MÃ SỐ THUẾ", "TÊN CÔNG TY", "MÃ SỐ THUẾ", "ĐƠN VỊ QUẢN LÝ THUẾ",
           "Văn bản", "Ngày", "Cơ quan ban hành", "Vụ án liên quan", "Mức rủi ro"]
    for j, h in enumerate(hdr, start=1):
        _set(ws1, 1, j, h)
    for i, (stt, mst, ten, vb) in enumerate(
        [(1, "0312816241", "CÔNG TY RƯỢU XYZ", "1396/YC-ANĐT-P6"),
         (2, "0313025323", "CÔNG TY GOLF", "1396/YC-ANĐT-P6")], start=2):
        _set(ws1, i, 1, stt); _set(ws1, i, 2, mst); _set(ws1, i, 3, ten); _set(ws1, i, 6, vb)

    ws2 = wb.create_sheet("Công văn 4532")
    _set(ws2, 3, 1, "Tên người nộp thuế"); _set(ws2, 3, 2, "Mã số thuế")
    _set(ws2, 3, 3, "Cơ quan thuế"); _set(ws2, 3, 4, "Mức rủi ro")
    _set(ws2, 4, 1, "CÔNG TY QUÀ"); _set(ws2, 4, 2, "0400111222"); _set(ws2, 4, 4, 3)

    ws3 = wb.create_sheet("524 DN")
    _set(ws3, 2, 2, "524 DN RỦI RO THEO Công văn số 1798/TCT-TTKT ngày 16/5/2023")
    _set(ws3, 4, 1, "Mã số thuế")
    _set(ws3, 5, 1, "0301224067")
    wb.save(os.path.join(OUT, "03_ds_rui_ro.xlsx"))


def make_einvoice():
    """File HĐĐT (mẫu): sheet DanhSach, tiêu đề dòng 4, cột 4/5/7/13/19."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DanhSach"
    _set(ws, 1, 1, "DANH SÁCH HÓA ĐƠN")
    # dòng 4: tiêu đề
    labels = {1: "STT", 2: "Ký hiệu mẫu số", 3: "Ký hiệu hóa đơn", 4: "Số hóa đơn",
              5: "Ngày lập", 7: "MST người bán", 8: "Tên người bán",
              13: "Tổng tiền thuế", 16: "Tổng tiền thanh toán", 19: "Trạng thái hóa đơn"}
    for c, t in labels.items():
        _set(ws, 4, c, t)
    rows = [
        # số, ngày,       mst,          thuế,   trạng thái
        ("123", "15/01/2022", "0301224067", 525000, "Hóa đơn mới"),
        ("45", "20/02/2022", "0312816241", 500000, "Hóa đơn đã bị thay thế"),
        ("46", "01/03/2022", "0313025323", 250000, "Hóa đơn mới"),   # lệch thuế 50k
        ("789", "10/04/2022", "0400111222", 300000, "Hóa đơn đã bị xóa bỏ/hủy bỏ"),
        # HĐ 790 không có -> không tìm thấy
    ]
    for i, (so, ngay, mst, thue, tt) in enumerate(rows, start=5):
        _set(ws, i, 4, so)
        _set(ws, i, 5, ngay)
        _set(ws, i, 7, mst)
        _set(ws, i, 13, thue)
        _set(ws, i, 19, tt)
    wb.save(os.path.join(OUT, "04_hoa_don_dien_tu.xlsx"))


if __name__ == "__main__":
    make_main()
    make_tms()
    make_risk()
    make_einvoice()
    print("Đã tạo file mẫu trong:", OUT)
